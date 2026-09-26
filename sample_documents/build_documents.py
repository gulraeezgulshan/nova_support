"""Render the Markdown policy sources in `sources/` to the PDF and DOCX files the
knowledge base ingests (the SRS requires both formats to be supported).

Usage (repo root): `uv run python -m sample_documents.build_documents`
"""

import html
import io
import re
from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).resolve().parent
SOURCES = HERE / "sources"
OUTPUT = HERE / "generated"

# Alternate formats so both parsers are exercised; FAQ stays Markdown (optional format).
FORMATS = {
    "DEL-POL-04": "pdf",
    "REF-POL-01": "docx",
    "WAR-POL-02": "pdf",
    "ESC-PRC-01": "docx",
    "RTG-RUL-01": "pdf",
    "SLA-AGR-01": "docx",
    "FAQ-GEN-01": "md",
}

CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.45; }
h1 { font-size: 20pt; font-weight: bold; margin-bottom: 8pt; }
h2 { font-size: 14pt; font-weight: bold; margin-top: 14pt; }
h3 { font-size: 12pt; font-weight: bold; margin-top: 10pt; }
p.meta { font-size: 9.5pt; color: #444; margin: 0; }
"""

HEADING = re.compile(r"^(#{1,3})\s+(.*)$")
META = re.compile(r"^[A-Z][A-Za-z ]+:\s")


def parse_markdown(text: str) -> list[tuple[str, str]]:
    """Return (kind, text) blocks: h1/h2/h3, meta or p."""
    blocks: list[tuple[str, str]] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        heading = HEADING.match(line)
        if heading:
            blocks.append((f"h{len(heading.group(1))}", heading.group(2)))
        elif META.match(line) and len(line) < 80:
            blocks.append(("meta", line))
        else:
            blocks.append(("p", line))
    return blocks


def to_pdf(blocks: list[tuple[str, str]], title: str) -> bytes:
    body = "".join(
        f'<p class="meta">{html.escape(t)}</p>'
        if kind == "meta"
        else f"<{kind}>{html.escape(t)}</{kind}>"
        for kind, t in blocks
    )
    story = pymupdf.Story(html=f"<body>{body}</body>", user_css=CSS)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    mediabox = pymupdf.paper_rect("a4")
    area = pymupdf.Rect(60, 60, mediabox.width - 60, mediabox.height - 60)  # page margins
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()

    pdf = pymupdf.open(stream=buffer.getvalue(), filetype="pdf")
    pdf.set_metadata({"title": title, "author": "VoltHaven Electronics (fictional)"})
    for index in range(pdf.page_count):
        footer = f"Page {index + 1} of {pdf.page_count}"
        pdf.load_page(index).insert_text((60, 815), footer, fontsize=8)
    data: bytes = pdf.tobytes(garbage=3, deflate=True)
    pdf.close()
    return data


def to_docx(blocks: list[tuple[str, str]], title: str) -> bytes:
    document = docx.Document()
    document.core_properties.title = title
    document.core_properties.author = "VoltHaven Electronics (fictional)"
    for kind, text in blocks:
        if kind == "h1":
            document.add_heading(text, level=0)
        elif kind in ("h2", "h3"):
            document.add_heading(text, level=int(kind[1]) - 1)
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def main() -> None:
    OUTPUT.mkdir(exist_ok=True)
    for source in sorted(SOURCES.glob("*.md")):
        doc_code = source.stem.split("_", 1)[0]
        fmt = FORMATS.get(doc_code, "pdf")
        text = source.read_text(encoding="utf-8")
        blocks = parse_markdown(text)
        title = next(t for kind, t in blocks if kind == "h1")
        target = OUTPUT / f"{source.stem}.{fmt}"
        if fmt == "pdf":
            target.write_bytes(to_pdf(blocks, title))
        elif fmt == "docx":
            target.write_bytes(to_docx(blocks, title))
        else:
            target.write_text(text, encoding="utf-8")
        print(f"{target.relative_to(HERE)}")


if __name__ == "__main__":
    main()
