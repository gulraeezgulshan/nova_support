"""Read a complaints file (.csv or .xlsx) into rows keyed by the known column names."""

import csv
import io
import zipfile

from openpyxl import Workbook, load_workbook

COLUMNS = [
    "customer_email",
    "customer_name",
    "title",
    "description",
    "order_ref",
    "channel",
    "received_at",
    "requested_resolution",
    "external_ref",
]
REQUIRED = ["customer_email", "title", "description"]
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
EXAMPLE = [
    "sara@example.com",
    "Sara Khan",
    "Late delivery",
    "My order was due last week and has still not arrived. Please send it or refund me.",
    "",
    "phone_callback",
    "2026-09-20T10:30:00",
    "Refund",
    "CALL-1001",
]


class ImportFileError(ValueError):
    pass


def escape_cell(value: str) -> str:
    """Stop spreadsheet apps from running a cell as a formula (CSV injection)."""
    return f"'{value}" if value[:1] in {"=", "+", "-", "@"} else value


def _table(filename: str, data: bytes) -> list[list[str]]:
    if not data:
        raise ImportFileError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ImportFileError("Files must be 5 MB or smaller.")
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ImportFileError("CSV files must be saved as UTF-8.") from exc
        return list(csv.reader(io.StringIO(text)))
    if name.endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise ImportFileError("The Excel file could not be read.") from exc
        sheet = book.worksheets[0]
        return [
            ["" if value is None else str(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    raise ImportFileError("Upload a CSV or .xlsx file.")


def read_rows(filename: str, data: bytes) -> tuple[list[dict[str, str]], list[str]]:
    """Rows as {column: value} (known columns, non-empty values) and the unknown columns."""
    table = [row for row in _table(filename, data) if any(cell.strip() for cell in row)]
    if not table:
        raise ImportFileError("The file is empty.")
    header = [h.strip().lower() for h in table[0]]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise ImportFileError(f"Missing required column(s): {', '.join(missing)}.")
    body = table[1:]
    if len(body) > MAX_ROWS:
        raise ImportFileError(f"At most {MAX_ROWS:,} rows per file.")
    unknown = [h for h in header if h and h not in COLUMNS]
    rows = [
        {
            column: row[i].strip()
            for i, column in enumerate(header)
            if column in COLUMNS and i < len(row) and row[i].strip()
        }
        for row in body
    ]
    return rows, unknown


def template_csv() -> bytes:
    buffer = io.StringIO()
    csv.writer(buffer).writerows([COLUMNS, EXAMPLE])
    return buffer.getvalue().encode("utf-8")


def template_xlsx() -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Complaints"
    sheet.append(COLUMNS)
    sheet.append(EXAMPLE)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
