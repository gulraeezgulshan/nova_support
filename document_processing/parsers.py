"""Extract structured text from PDF, DOCX, TXT and Markdown files (SRS Step 5).

Output is a list of sections (heading + numbered section + paragraphs), and every
paragraph remembers the page it came from, so chunks can cite page numbers.
DOCX files have no fixed pagination, so their page numbers are None.
"""

import io
import re
import unicodedata
from collections import Counter
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field

import docx
import pymupdf
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph as DocxParagraph

SECTION_NUMBER = re.compile(r"^(?:section\s+)?(\d+(?:\.\d+)*)[.)]?\s+(\S.*)$", re.IGNORECASE)
PAGE_FOOTER = re.compile(r"^(page\s+)?\d+(\s*(of|/)\s*\d+)?$", re.IGNORECASE)
MARKDOWN_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*$")
# "Document ID: DEL-POL-04" style lines are kept as their own paragraph, never merged.
LABEL_LINE = re.compile(r"^[A-Z][A-Za-z /]{1,30}:\s+\S")


def normalise_text(text: str) -> str:
    """NFKC turns ligatures (e.g. U+FB00 'ff') and full-width forms into plain characters."""
    return " ".join(unicodedata.normalize("NFKC", text).split())


@dataclass
class Paragraph:
    text: str
    page: int | None


@dataclass
class Section:
    heading: str | None
    section_number: str | None
    level: int
    paragraphs: list[Paragraph] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.paragraphs)

    @property
    def page_start(self) -> int | None:
        pages = [p.page for p in self.paragraphs if p.page is not None]
        return min(pages) if pages else None

    @property
    def page_end(self) -> int | None:
        pages = [p.page for p in self.paragraphs if p.page is not None]
        return max(pages) if pages else None


@dataclass
class ParsedDocument:
    title: str | None
    page_count: int | None
    sections: list[Section]

    @property
    def full_text(self) -> str:
        parts: list[str] = []
        for section in self.sections:
            if section.heading:
                parts.append(section.heading)
            parts.append(section.text)
        return "\n".join(p for p in parts if p)


@dataclass
class _Block:
    """Intermediate unit shared by all parsers before sections are assembled."""

    text: str
    page: int | None
    heading_level: int | None = None  # None = body text


class ParseError(Exception):
    pass


def parse_document(data: bytes, file_type: str) -> ParsedDocument:
    title: str | None
    page_count: int | None
    blocks: list[_Block]
    if file_type == "pdf":
        title, page_count, blocks = _pdf_blocks(data)
    elif file_type == "docx":
        title, page_count, blocks = _docx_blocks(data)
    elif file_type in ("txt", "md"):
        title, page_count, blocks = None, None, list(_text_blocks(data.decode("utf-8")))
    else:
        raise ParseError(f"Unsupported file type: {file_type}")

    sections = _assemble_sections(blocks)
    if title is None:
        title = next((s.heading for s in sections if s.heading), None)
    return ParsedDocument(title=title, page_count=page_count, sections=sections)


# --------------------------------------------------------------------------- PDF


def _pdf_blocks(data: bytes) -> tuple[str | None, int, list[_Block]]:
    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # PyMuPDF raises several internal exception types
        raise ParseError(f"Could not open PDF: {exc}") from exc

    with pdf:
        if pdf.needs_pass:
            raise ParseError("PDF is password protected.")
        lines: list[tuple[int, str, float, bool]] = []  # page, text, font size, bold
        for page_index in range(1, pdf.page_count + 1):
            page = pdf.load_page(page_index - 1)
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 0:  # 0 = text, 1 = image
                    continue
                for line in block["lines"]:
                    spans = [s for s in line["spans"] if s["text"].strip()]
                    if not spans:
                        continue
                    text = normalise_text(" ".join(s["text"] for s in spans))
                    size = max(s["size"] for s in spans)
                    bold = all(s["flags"] & 16 or "bold" in s["font"].lower() for s in spans)
                    lines.append((page_index, text, size, bold))
                # Block boundary marks a paragraph break.
                lines.append((page_index, "", 0.0, False))
        metadata_title = (pdf.metadata or {}).get("title") or None
        page_count = pdf.page_count

    body_size = _dominant_font_size(lines)
    blocks: list[_Block] = []
    buffer: list[str] = []
    buffer_page: int | None = None

    def flush() -> None:
        nonlocal buffer, buffer_page
        if buffer:
            blocks.append(_Block(" ".join(buffer), buffer_page))
        buffer, buffer_page = [], None

    for page, text, size, bold in lines:
        if not text:
            flush()
            continue
        if PAGE_FOOTER.match(text):
            continue
        level = _pdf_heading_level(text, size, bold, body_size)
        if level is not None:
            flush()
            blocks.append(_Block(text, page, level))
        elif LABEL_LINE.match(text):
            flush()
            blocks.append(_Block(text, page))
        else:
            if buffer_page is None:
                buffer_page = page
            buffer.append(text)
    flush()
    return metadata_title, page_count, blocks


def _dominant_font_size(lines: Iterable[tuple[int, str, float, bool]]) -> float:
    weights: Counter[float] = Counter()
    for _, text, size, _ in lines:
        if text:
            weights[round(size, 1)] += len(text)
    return weights.most_common(1)[0][0] if weights else 11.0


def _pdf_heading_level(text: str, size: float, bold: bool, body_size: float) -> int | None:
    if len(text) > 120 or text.endswith((".", ",", ";")):
        return None
    numbered = SECTION_NUMBER.match(text)
    if size >= body_size * 1.5:
        return 1
    if size >= body_size * 1.15:
        return _level_from_number(numbered.group(1)) if numbered else 2
    if bold and (numbered or len(text) <= 80):
        return _level_from_number(numbered.group(1)) if numbered else 3
    return None


def _level_from_number(number: str) -> int:
    return min(number.count(".") + 2, 6)  # "5" -> 2, "5.2" -> 3


# -------------------------------------------------------------------------- DOCX


def _docx_blocks(data: bytes) -> tuple[str | None, None, list[_Block]]:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several exception types
        raise ParseError(f"Could not open DOCX: {exc}") from exc

    title = (document.core_properties.title or "").strip() or None
    blocks: list[_Block] = []
    for item in _iter_docx_body(document):
        if isinstance(item, Table):
            for row in item.rows:
                cells = [normalise_text(c.text) for c in row.cells]
                row_text = " | ".join(dict.fromkeys(c for c in cells if c))  # merged cells repeat
                if row_text:
                    blocks.append(_Block(row_text, None))
            continue
        text = normalise_text(item.text)
        if not text:
            continue
        style = (item.style.name if item.style is not None else "") or ""
        level = _docx_heading_level(style, item, text)
        if style == "Title" and title is None:
            title = text
        blocks.append(_Block(text, None, level))
    return title, None, blocks


def _iter_docx_body(document: DocxDocument) -> Iterator[DocxParagraph | Table]:
    """Yield paragraphs and tables in document order."""
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield DocxParagraph(child, document)
        elif tag == "tbl":
            yield Table(child, document)


def _docx_heading_level(style: str, paragraph: DocxParagraph, text: str) -> int | None:
    if style == "Title":
        return 1
    match = re.match(r"^Heading\s*(\d)$", style)
    if match:
        return min(int(match.group(1)) + 1, 6)
    runs = [r for r in paragraph.runs if r.text.strip()]
    numbered = SECTION_NUMBER.match(text)
    if runs and all(r.bold for r in runs) and len(text) <= 100 and not text.endswith("."):
        return _level_from_number(numbered.group(1)) if numbered else 3
    return None


# ---------------------------------------------------------------------- TXT / MD


def _text_blocks(text: str) -> Iterator[_Block]:
    buffer: list[str] = []
    for raw_line in text.splitlines():
        line = normalise_text(raw_line)
        heading = MARKDOWN_HEADING.match(line)
        numbered = SECTION_NUMBER.match(line)
        is_numbered_heading = bool(numbered) and len(line) <= 90 and not line.endswith(".")
        if not line or heading or is_numbered_heading:
            if buffer:
                yield _Block(" ".join(buffer), None)
                buffer = []
            if heading:
                yield _Block(heading.group(2), None, len(heading.group(1)))
            elif is_numbered_heading and numbered:
                yield _Block(line, None, _level_from_number(numbered.group(1)))
            continue
        if LABEL_LINE.match(line):
            if buffer:
                yield _Block(" ".join(buffer), None)
                buffer = []
            yield _Block(line, None)
            continue
        buffer.append(line.lstrip("-*• ").strip() if line[:2] in ("- ", "* ", "• ") else line)
    if buffer:
        yield _Block(" ".join(buffer), None)


# ---------------------------------------------------------------------- sections


def _assemble_sections(blocks: list[_Block]) -> list[Section]:
    sections: list[Section] = []
    current = Section(heading=None, section_number=None, level=0)
    for block in blocks:
        if block.heading_level is not None:
            if current.paragraphs:
                sections.append(current)
            numbered = SECTION_NUMBER.match(block.text)
            current = Section(
                heading=block.text[:255],
                section_number=numbered.group(1) if numbered else None,
                level=block.heading_level,
            )
        else:
            current.paragraphs.append(Paragraph(block.text, block.page))
    if current.paragraphs:
        sections.append(current)
    return sections
