"""Report exporters (SRS Step 68): CSV (Excel-compatible, UTF-8 with BOM), Excel (.xlsx,
one sheet per table) and PDF."""

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    LongTable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)

from src.analytics.reports import Report, Table
from src.core.domain import organization


def cell(value: Any) -> str:
    """Display form of a value in any export format."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:g}"
    if isinstance(value, (list, tuple, set)):
        return ", ".join(cell(v) for v in value)
    return str(value)


def _header(report: Report) -> list[list[str]]:
    lines = [
        [report.title],
        [organization().get("name", "")],
        [f"Generated {report.generated_at:%Y-%m-%d %H:%M} UTC"],
    ]
    if report.filters:
        lines.append(["Filters: " + "; ".join(f"{k}: {v}" for k, v in report.filters)])
    return lines


# --- CSV ---------------------------------------------------------------------------


def to_csv(report: Report) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows(_header(report))
    if report.summary:
        writer.writerow([])
        writer.writerow(["Metric", "Value"])
        writer.writerows([label, cell(value)] for label, value in report.summary)
    for table in report.tables:
        writer.writerow([])
        writer.writerow([table.title])
        writer.writerow([header for _, header in table.columns])
        writer.writerows([cell(row.get(key)) for key, _ in table.columns] for row in table.rows)
    return ("﻿" + buffer.getvalue()).encode("utf-8")  # BOM: Excel opens it as UTF-8


# --- Excel -------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet_name(title: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Sheet"
    name, n = base, 2
    while name in used:
        suffix = f" ({n})"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(name)
    return name


def _xlsx_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)  # Excel has no time zones; values are UTC
    if isinstance(value, (int, float, date)) and not isinstance(value, bool):
        return value
    return cell(value)


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    used = {"Summary"}
    summary.title = "Summary"
    for line in _header(report):
        summary.append(line)
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    summary.append(["Metric", "Value"])
    for c in summary[summary.max_row]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for label, value in report.summary:
        summary.append([label, _xlsx_value(value)])
    summary.column_dimensions["A"].width = 45
    summary.column_dimensions["B"].width = 18

    for table in report.tables:
        ws = wb.create_sheet(_sheet_name(table.title, used))
        ws.append([header for _, header in table.columns])
        for c in ws[1]:
            c.fill, c.font = HEADER_FILL, HEADER_FONT
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in table.rows:
            ws.append([_xlsx_value(row.get(key)) for key, _ in table.columns])
        ws.freeze_panes = "A2"
        if table.rows:
            ws.auto_filter.ref = ws.dimensions
        for index, (key, header) in enumerate(table.columns, start=1):
            longest = max([len(header), *(len(cell(r.get(key))) for r in table.rows[:500])])
            ws.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)
            if isinstance(table.rows[0].get(key) if table.rows else None, datetime):
                for (c,) in ws.iter_rows(min_row=2, min_col=index, max_col=index):
                    c.number_format = "yyyy-mm-dd hh:mm"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- PDF ---------------------------------------------------------------------------

STYLES = getSampleStyleSheet()
CELL = ParagraphStyle("cell", parent=STYLES["BodyText"], fontSize=7, leading=8.5)
HEAD = ParagraphStyle("head", parent=CELL, textColor=colors.white, fontName="Helvetica-Bold")
TABLE_STYLE = TableStyle(
    [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C2CC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F8")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
)
PDF_ROW_LIMIT = 1000  # a PDF is for reading; the full data is in CSV and Excel


def _pdf_table(table: Table, width: float) -> list[Any]:
    rows = table.rows[:PDF_ROW_LIMIT]
    data = [[Paragraph(escape(h), HEAD) for _, h in table.columns]]
    data += [[Paragraph(escape(cell(r.get(k))), CELL) for k, _ in table.columns] for r in rows]
    weights = [
        min(max(len(h), *(len(cell(r.get(k))) for r in rows[:200]), 6), 40)
        for k, h in table.columns
    ]
    total = sum(weights)
    grid = LongTable(data, colWidths=[width * w / total for w in weights], repeatRows=1)
    grid.setStyle(TABLE_STYLE)
    flow: list[Any] = [Paragraph(escape(table.title), STYLES["Heading3"])]
    if not table.rows:
        return [*flow, Paragraph("No data.", STYLES["BodyText"]), Spacer(1, 4 * mm)]
    if len(table.rows) > PDF_ROW_LIMIT:
        flow.append(
            Paragraph(
                f"First {PDF_ROW_LIMIT} of {len(table.rows)} rows; the CSV and "
                "Excel exports contain all rows.",
                CELL,
            )
        )
    if len(rows) < 25:  # keep short tables on one page with their heading
        return [KeepTogether([*flow, grid]), Spacer(1, 5 * mm)]
    return [*flow, grid, Spacer(1, 5 * mm)]


def to_pdf(report: Report) -> bytes:
    buffer = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(
        buffer,
        pagesize=page,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=report.title,
        author=str(organization().get("name", "SupportNova")),
    )
    width = page[0] - 24 * mm
    story: list[Any] = [Paragraph(escape(report.title), STYLES["Title"])]
    for (line,) in _header(report)[1:]:
        story.append(Paragraph(escape(line), STYLES["BodyText"]))
    story.append(Paragraph(escape(report.description), STYLES["Italic"]))
    story.append(Spacer(1, 4 * mm))
    if report.summary:
        story += _pdf_table(
            Table(
                "Summary",
                [("label", "Metric"), ("value", "Value")],
                [{"label": k, "value": v} for k, v in report.summary],
            ),
            width * 0.6,
        )
    for table in report.tables:
        story += _pdf_table(table, width)

    def footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(page[0] - 12 * mm, 6 * mm, f"{report.title} · page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


@dataclass(frozen=True)
class Exporter:
    extension: str
    media_type: str
    render: Callable[[Report], bytes]


EXPORTERS: dict[str, Exporter] = {
    "csv": Exporter("csv", "text/csv; charset=utf-8", to_csv),
    "xlsx": Exporter(
        "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", to_xlsx
    ),
    "pdf": Exporter("pdf", "application/pdf", to_pdf),
}
