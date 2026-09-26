"""Report exports: CSV (Excel-compatible), Excel and PDF."""

import csv
import io
from datetime import UTC, datetime

import pymupdf
from openpyxl import load_workbook

from src.analytics.export import EXPORTERS, _sheet_name, cell, to_csv, to_pdf, to_xlsx
from src.analytics.reports import REPORTS, Report, Table

GENERATED = datetime(2026, 9, 26, 10, 30, tzinfo=UTC)


def sample_report() -> Report:
    rows = [
        {"ref": "CMP-000001", "created": GENERATED, "level": 3, "flag": True, "note": "Café, ok"},
        {"ref": "CMP-000002", "created": GENERATED, "level": 0, "flag": False, "note": None},
    ]
    return Report(
        code="sample",
        title="Sample Report",
        description="For tests.",
        generated_at=GENERATED,
        filters=[("Category", "DELIVERY")],
        summary=[("Total complaints", 2), ("SLA compliance %", 87.5)],
        tables=[
            Table(
                "Complaints",
                [
                    ("ref", "Complaint"),
                    ("created", "Submitted"),
                    ("level", "Escalation"),
                    ("flag", "Flag"),
                    ("note", "Note"),
                ],
                rows,
            ),
            Table("Empty: section / with [odd] name?", [("x", "X")], []),
        ],
    )


def test_cell_formats_values_for_people() -> None:
    assert cell(None) == ""
    assert cell(True) == "Yes"
    assert cell(GENERATED) == "2026-09-26 10:30"
    assert cell(87.50) == "87.5"
    assert cell(["a", "b"]) == "a, b"


def test_csv_has_bom_title_filters_summary_and_every_table() -> None:
    data = to_csv(sample_report())
    assert data.startswith(b"\xef\xbb\xbf")  # Excel opens UTF-8 correctly
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0] == ["Sample Report"]
    assert ["Filters: Category: DELIVERY"] in rows
    assert ["Total complaints", "2"] in rows
    header = rows.index(["Complaint", "Submitted", "Escalation", "Flag", "Note"])
    assert rows[header + 1] == ["CMP-000001", "2026-09-26 10:30", "3", "Yes", "Café, ok"]
    assert ["Empty: section / with [odd] name?"] in rows


def test_xlsx_has_a_summary_and_one_sheet_per_table() -> None:
    wb = load_workbook(io.BytesIO(to_xlsx(sample_report())))
    assert wb.sheetnames == ["Summary", "Complaints", "Empty section  with odd name"]
    sheet = wb["Complaints"]
    assert [c.value for c in sheet[1]] == ["Complaint", "Submitted", "Escalation", "Flag", "Note"]
    assert sheet["C2"].value == 3  # numbers stay numbers
    assert sheet["B2"].value == datetime(2026, 9, 26, 10, 30)
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row[0]}
    assert summary["SLA compliance %"] == 87.5


def test_sheet_names_are_valid_and_unique() -> None:
    used: set[str] = set()
    names = [_sheet_name("A very long table title that exceeds the limit", used) for _ in range(3)]
    assert len(set(names)) == 3 and all(len(n) <= 31 for n in names)


def test_pdf_contains_the_report() -> None:
    data = to_pdf(sample_report())
    assert data.startswith(b"%PDF")
    with pymupdf.open(stream=data, filetype="pdf") as pdf:  # type: ignore[no-untyped-call]
        text = "".join(pdf.load_page(i).get_text() for i in range(pdf.page_count))
    assert "Sample Report" in text and "CMP-000002" in text and "No data." in text


def test_every_report_has_a_spec_and_every_format_an_exporter() -> None:
    assert {
        "complaint_intelligence",
        "department_performance",
        "escalations",
        "sla_status",
        "policy_usage",
        "resolution_compliance",
        "genai_python_comparison",
        "manual_reviews",
        "complaint_analysis",
    } <= set(REPORTS)
    assert set(EXPORTERS) == {"csv", "xlsx", "pdf"}
