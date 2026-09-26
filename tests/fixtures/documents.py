"""Build small PDF / DOCX files in memory for tests."""

import io

import docx
import pymupdf

# (text, font size, bold). A None entry starts a new page.
PdfLine = tuple[str, float, bool] | None


def make_pdf(lines: list[PdfLine]) -> bytes:
    pdf = pymupdf.open()
    page = pdf.new_page()
    y = 72.0
    for line in lines:
        if line is None:
            page.insert_text((300, 800), "Page 1 of 2", fontsize=8, fontname="helv")
            page = pdf.new_page()
            y = 72.0
            continue
        text, size, bold = line
        page.insert_text((72, y), text, fontsize=size, fontname="hebo" if bold else "helv")
        y += size * 2.2  # blank space between lines -> separate text blocks
    data: bytes = pdf.tobytes()
    pdf.close()
    return data


def make_docx(
    header: dict[str, str], sections: list[tuple[str, list[str]]], title: str = "Delivery Policy"
) -> bytes:
    document = docx.Document()
    document.add_heading(title, level=0)
    for label, value in header.items():
        document.add_paragraph(f"{label}: {value}")
    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


DELIVERY_HEADER = {
    "Document ID": "DEL-POL-04",
    "Document Type": "policy",
    "Version": "1.0",
    "Effective Date": "2026-01-01",
}

DELIVERY_SECTIONS = [
    ("1 Purpose", ["This policy defines delivery commitments for VoltHaven orders."]),
    (
        "5.2 Late Delivery Compensation",
        [
            "If a standard order arrives more than 5 business days late, the customer is "
            "eligible for a 10% store credit on the order value.",
            "Compensation is not offered for delays caused by incorrect addresses.",
        ],
    ),
]
