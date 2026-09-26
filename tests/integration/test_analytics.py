"""Analytics, trends, SLA scanning, dashboards, reports and exports against the database."""

import hashlib
import io
import uuid
from collections.abc import Callable
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from complaint_processing.sla import scan
from database.models import (
    Complaint,
    ComplaintEvent,
    Customer,
    CustomerType,
    Order,
    Role,
    User,
)
from database.session import sync_session
from python_validation.pipeline import run_validation
from src.analytics import metrics
from src.analytics.filters import Filters
from src.analytics.trends import detect_trends

pytestmark = pytest.mark.db

NOW = datetime.now(UTC)


def add_complaints(specs: list[dict[str, Any]]) -> None:
    """Insert complaints directly with the given fields (defaults: an open delivery case)."""
    with sync_session() as db:
        customer = Customer(
            customer_ref="CUST-910001", full_name="Dana", customer_type=CustomerType.STANDARD
        )
        db.add(customer)
        db.flush()
        order = Order(
            order_ref="ORD-910001",
            customer_id=customer.id,
            product_name="AeroBook 14",
            product_category="LAPTOP",
            amount=Decimal("899.00"),
            order_date=date(2026, 9, 1),
            committed_delivery_date=date(2026, 9, 8),
        )
        db.add(order)
        db.flush()
        for i, spec in enumerate(specs, start=1):
            text = f"complaint {i} {uuid.uuid4()}"
            fields: dict[str, Any] = {
                "complaint_ref": f"CMP-{900000 + i}",
                "customer_id": customer.id,
                "order_id": order.id,
                "title": f"Complaint {i}",
                "description": text,
                "normalized_text": text,
                "content_hash": hashlib.sha256(text.encode()).hexdigest(),
                "created_at": NOW - timedelta(days=1),
                "category_code": "DELIVERY",
                "subcategory_code": "DELAYED_DELIVERY",
                "department_code": "LOGISTICS",
                "priority": "P2",
                "sentiment": "negative",
                "status": "assigned",
                "sla_status": "pending",
                **spec,
            }
            db.add(Complaint(**fields))


@pytest.fixture
def portfolio(clean_db: None) -> None:
    add_complaints(
        [
            {"escalation_level": 2, "status": "escalated", "priority": "P1"},
            {
                "status": "resolved",
                "created_at": NOW - timedelta(days=3),
                "resolved_at": NOW - timedelta(days=2),
                "sla_status": "met",
            },
            {
                "category_code": "BILLING",
                "subcategory_code": "DUPLICATE_CHARGE",
                "department_code": "BILLING",
                "sentiment": "angry",
                "sla_status": "breached",
            },
            {"created_at": NOW - timedelta(days=10), "signals": {"repeat_contact": ["again"]}},
            {
                "category_code": None,
                "subcategory_code": None,
                "department_code": None,
                "priority": None,
                "status": "new",
            },
        ]
    )


def test_overview_and_distributions(portfolio: None) -> None:
    with sync_session() as db:
        o = metrics.overview(db, Filters())
        assert (o["total"], o["open"], o["resolved"], o["escalated"]) == (5, 4, 1, 1)
        assert (o["sla_breached"], o["sla_met"], o["repeat"], o["unclassified"]) == (1, 1, 1, 1)
        assert o["sla_compliance_pct"] == 100.0 and o["avg_resolution_hours"] == 24.0
        categories = {d["key"]: d["count"] for d in metrics.distribution(db, Filters(), "category")}
        assert categories == {"DELIVERY": 3, "BILLING": 1, "UNCLASSIFIED": 1}
        labels = {d["key"]: d["label"] for d in metrics.distribution(db, Filters(), "category")}
        assert labels["DELIVERY"] != "DELIVERY"  # display names come from the taxonomy
        assert metrics.overview(db, Filters(category="BILLING"))["total"] == 1
        recent = Filters(date_from=(NOW - timedelta(days=5)).date())
        assert metrics.overview(db, recent)["total"] == 4
        departments = {d["department"]: d for d in metrics.department_performance(db, Filters())}
        assert departments["LOGISTICS"]["escalated"] == 1
        assert departments["BILLING"]["breached"] == 1


def test_trends_detect_rising_categories_recurring_issues_and_spikes(clean_db: None) -> None:
    this_week = [{"created_at": NOW - timedelta(days=d)} for d in (1, 2, 3, 4)]
    escalated = [{"created_at": NOW - timedelta(days=d), "escalation_level": 1} for d in (1, 2, 3)]
    last_week = [
        {
            "created_at": NOW - timedelta(days=9),
            "category_code": "BILLING",
            "subcategory_code": "DUPLICATE_CHARGE",
        }
    ]
    add_complaints(this_week + escalated + last_week)
    with sync_session() as db:
        trends = {(t.kind, t.key): t for t in detect_trends(db)}
    delivery = trends[("rising_category", "DELIVERY")]
    assert (delivery.current, delivery.previous) == (7, 0)
    assert "delivery" in delivery.label.lower()
    assert ("recurring_product_issue", "LAPTOP|DELAYED_DELIVERY") in trends
    assert trends[("escalation_spike", "all")].current == 3
    assert ("rising_category", "BILLING") not in trends


def test_sla_scan_flags_a_breach_once(clean_db: None) -> None:
    add_complaints(
        [
            {"priority": "P0", "created_at": NOW - timedelta(hours=5)},
            {"priority": "P3", "created_at": NOW - timedelta(hours=1)},
        ]
    )
    with sync_session() as db:
        first = scan(db)
        assert first.newly_breached == ["CMP-900001"]
        assert first.counts == {"breached": 1, "on_track": 1}
        assert scan(db).changed == 0
        events = db.scalars(select(ComplaintEvent).where(ComplaintEvent.event_type.like("sla_%")))
        [event] = events.all()
        assert event.event_type == "sla_breached" and not event.customer_visible


def test_validation_starts_the_sla_clock(complaint_id: uuid.UUID) -> None:
    with sync_session() as db:
        run_validation(db, complaint_id)  # no GenAI analysis: Python only
        complaint = db.get(Complaint, complaint_id)
        assert complaint is not None and complaint.priority
        assert complaint.resolution_due_at and complaint.first_response_due_at
        assert complaint.sla_status in {"on_track", "at_risk"}


# --- API ---------------------------------------------------------------------------


@pytest.fixture
def reviewer(auth_headers: Callable[[Role], dict[str, str]]) -> dict[str, str]:
    return auth_headers(Role.REVIEWER)


async def test_admin_dashboard_is_for_reviewers_managers_and_admins(
    client: httpx.AsyncClient,
    portfolio: None,
    reviewer: dict[str, str],
    auth_headers: Callable[[Role], dict[str, str]],
) -> None:
    denied = await client.get("/api/v1/dashboard/admin", headers=auth_headers(Role.AGENT))
    assert denied.status_code == 403
    body = (await client.get("/api/v1/dashboard/admin", headers=reviewer)).json()
    assert body["overview"]["total"] == 5
    assert {"category", "department", "priority", "sla_status"} <= set(body["distributions"])
    assert [r["complaint_ref"] for r in body["sla_risks"]] == ["CMP-900003"]
    filtered = await client.get(
        "/api/v1/dashboard/admin", params={"department": "BILLING"}, headers=reviewer
    )
    assert filtered.json()["overview"]["total"] == 1


async def test_analytics_endpoint(
    client: httpx.AsyncClient, portfolio: None, reviewer: dict[str, str]
) -> None:
    response = await client.get("/api/v1/analytics", params={"bucket": "week"}, headers=reviewer)
    assert response.status_code == 200, response.text
    body = response.json()
    assert sum(point["total"] for point in body["volume"]) == 5
    assert {"product", "urgency", "sentiment", "escalation_level"} <= set(body["distributions"])
    assert body["distributions"]["product"][0]["key"] == "LAPTOP"


async def test_agent_dashboard_shows_the_department_queue_with_warnings(
    client: httpx.AsyncClient, portfolio: None, auth_headers: Callable[[Role], dict[str, str]]
) -> None:
    agent = auth_headers(Role.AGENT)
    body = (
        await client.get("/api/v1/dashboard/agent", params={"department": "BILLING"}, headers=agent)
    ).json()
    assert body["department"] == "BILLING"
    [item] = body["items"]
    assert item["complaint_ref"] == "CMP-900003" and "SLA breached" in item["escalation_warnings"]
    everyone = (await client.get("/api/v1/dashboard/agent", headers=agent)).json()
    assert len(everyone["items"]) == 4  # every open complaint when the agent has no department
    assert everyone["items"][0]["sla_status"] == "breached"  # most urgent first


async def test_reports_preview_and_exports(
    client: httpx.AsyncClient, portfolio: None, reviewer: dict[str, str]
) -> None:
    specs = (await client.get("/api/v1/reports", headers=reviewer)).json()
    assert len(specs) == 9
    preview = (await client.get("/api/v1/reports/complaint_intelligence", headers=reviewer)).json()
    titles = [t["title"] for t in preview["tables"]]
    assert "Category distribution" in titles and "Manual-review cases" in titles
    assert dict((s["label"], s["value"]) for s in preview["summary"])["Total complaints"] == "5"

    for fmt, media in [("csv", "text/csv"), ("xlsx", "spreadsheetml"), ("pdf", "application/pdf")]:
        response = await client.get(
            "/api/v1/reports/department_performance/export",
            params={"format": fmt, "category": "DELIVERY"},
            headers=reviewer,
        )
        assert response.status_code == 200, response.text
        assert media in response.headers["content-type"]
        assert f'.{fmt}"' in response.headers["content-disposition"]
        if fmt == "xlsx":
            workbook = load_workbook(io.BytesIO(response.content))
            assert "Department performance" in workbook.sheetnames
            rows = list(workbook["Department performance"].iter_rows(values_only=True))
            assert [r[0] for r in rows[1:]] == ["Logistics"] or len(rows) == 2
    missing = await client.get("/api/v1/reports/nope", headers=reviewer)
    assert missing.status_code == 404


async def test_complaint_search_filters(
    client: httpx.AsyncClient, portfolio: None, reviewer: dict[str, str]
) -> None:
    async def refs(**params: Any) -> list[str]:
        response = await client.get("/api/v1/complaints", params=params, headers=reviewer)
        assert response.status_code == 200, response.text
        return sorted(item["complaint_ref"] for item in response.json()["items"])

    assert await refs(escalated=True) == ["CMP-900001"]
    assert await refs(sla_status="breached") == ["CMP-900003"]
    assert await refs(sentiment="angry") == ["CMP-900003"]
    assert await refs(date_to=(NOW - timedelta(days=5)).date().isoformat()) == ["CMP-900004"]


async def test_customer_list_shows_latest_update(
    client: httpx.AsyncClient, create_user: Callable[..., User], make_token: Callable[..., str]
) -> None:
    user = create_user(Role.CUSTOMER)
    headers = {"Authorization": f"Bearer {make_token(user.clerk_user_id)}"}
    payload = {
        "title": "Speaker broken",
        "description": "My HomeHub speaker stopped working after three days of use.",
    }
    assert (await client.post("/api/v1/complaints", json=payload, headers=headers)).is_success
    [item] = (await client.get("/api/v1/complaints", headers=headers)).json()["items"]
    assert item["latest_update"] and item["latest_update_at"]
    assert item["sla_status"] is None  # SLA details are staff-only


async def test_vocabulary_comes_from_configuration(
    client: httpx.AsyncClient, auth_headers: Callable[[Role], dict[str, str]]
) -> None:
    body = (
        await client.get("/api/v1/taxonomy/vocabulary", headers=auth_headers(Role.AGENT))
    ).json()
    assert "Strongly Negative" in body["sentiments"] and "P0" in body["priorities"]
    assert "web_form" in body["channels"]
    assert {"code": "ISSUE_FULL_REFUND", "kind": "remedy"}.items() <= next(
        a for a in body["actions"] if a["code"] == "ISSUE_FULL_REFUND"
    ).items()
